"""Offline smoke check; outputs live only in a temporary directory."""
import os
from pathlib import Path
from tempfile import TemporaryDirectory
cache = Path("/workspace/.cache/paper-tools")
(cache / "matplotlib").mkdir(parents=True, exist_ok=True)
(cache / "xdg").mkdir(parents=True, exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(cache / "matplotlib"))
os.environ.setdefault("XDG_CACHE_HOME", str(cache / "xdg"))
import numpy as np
import scipy
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document

with TemporaryDirectory(prefix="paper-workflow-") as directory:
    root = Path(directory)
    doc = Document()
    doc.add_heading("Gear research workflow check", 0)
    doc.add_paragraph("Mesh frequency and contact response: tool smoke check only.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Time (s)"
    table.cell(0, 1).text = "Displacement (m)"
    path = root / "check.docx"
    doc.save(path)
    loaded = Document(path)
    assert loaded.paragraphs[0].text == "Gear research workflow check"
    assert loaded.tables[0].cell(0, 1).text == "Displacement (m)"
    t = np.linspace(0, 1, 100)
    fig, ax = plt.subplots()
    ax.plot(t, np.sin(2 * np.pi * t))
    ax.set(xlabel="Time (s)", ylabel="Synthetic signal (dimensionless)")
    for extension in ("svg", "pdf", "png"):
        output = root / f"check.{extension}"
        fig.savefig(output, dpi=300)
        assert output.stat().st_size > 100
    plt.close(fig)
print("PASS: DOCX paragraph/table round-trip and SVG/PDF/PNG scientific plotting")
